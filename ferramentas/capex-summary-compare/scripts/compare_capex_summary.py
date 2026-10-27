#!/usr/bin/env python3
"""Compara a aba 'CapEx Summary' de duas versoes do PLA-FOR-010 (Cons = base, PROD = atual)
e gera um dashboard HTML autocontido + JSON com as diferencas.

Uso:
  python3 -I compare_capex_summary.py --folder <pasta> --template <template.html> --out <saida.html>
  python3 -I compare_capex_summary.py --old <cons.xlsx> --new <prod.xlsx> --template ... --out ...
Opcoes: --tol 1 (tolerancia em USD para considerar valor alterado)
"""
import argparse, datetime as dt, glob, json, os, re, sys
import warnings
warnings.filterwarnings("ignore")
import openpyxl

SHEET = "CapEx Summary"
# Projetos de ajuste (ex.: "Ajuste Brasil", "Ajuste Mexico") ficam sempre fora da analise, nas duas versoes.
EXCLUDE_PROJECT = re.compile(r"^\s*ajuste\b", re.I)
METRICS = ["Total Budget", "EFC", "EFC Non-Spec", "EFC Spec", "Spend to Date", "Remaining Capex", "% Procured"]
ATTRS = ["Capex categorization", "Type", "Fitout", "Support Debt Covenants", "Other", "Commit?",
         "BBNB Capex", "Country"]
TOTAL_METRICS = ["Total Budget", "EFC", "Spend to Date", "Remaining Capex"]
COUNTRY_FIX = {"Brasil": "Brazil", "Mexico": "Mexico", "México": "Mexico"}


# ---------------------------------------------------------------- localizar arquivos
def name_date(path):
    m = re.search(r"(\d{2})-(\d{2})-(\d{2,4})", os.path.basename(path))
    if not m:
        return dt.date.min
    d, mth, y = map(int, m.groups())
    y = y + 2000 if y < 100 else y
    try:
        return dt.date(y, mth, d)
    except ValueError:
        return dt.date.min


def pick_files(folder):
    files = [f for f in glob.glob(os.path.join(folder, "*.xls*"))
             if not os.path.basename(f).startswith("~$")]
    cons = [f for f in files if re.search(r"\bcons", os.path.basename(f), re.I)]
    prod = [f for f in files if re.search(r"\bprod\b", os.path.basename(f), re.I)]
    if not cons or not prod:
        sys.exit(f"ERRO: preciso de um arquivo com 'Cons' e outro com 'PROD' no nome em {folder}. "
                 f"Encontrados: {[os.path.basename(f) for f in files]}")
    key = lambda f: (name_date(f), os.path.getmtime(f))
    return max(cons, key=key), max(prod, key=key)


# ---------------------------------------------------------------- leitura
def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None


def txt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def read_summary(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        sys.exit(f"ERRO: aba '{SHEET}' nao encontrada em {os.path.basename(path)}")
    rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]
    wb.close()

    # linha de cabecalho: A='Site' e B='Phase/Project'
    hi = next((i for i, r in enumerate(rows[:30])
               if len(r) > 1 and txt(r[0]) == "Site" and txt(r[1]) == "Phase/Project"), None)
    if hi is None:
        sys.exit(f"ERRO: cabecalho 'Site | Phase/Project' nao encontrado em {os.path.basename(path)}")
    hdr = rows[hi]
    col = {}
    for j, h in enumerate(hdr):
        k = txt(h)
        if k and not isinstance(h, dt.datetime) and k not in col:
            col[k] = j
    col.setdefault("Fitout", col.get("Fitout"))
    missing = [m for m in ["Total Budget", "EFC", "Spend to Date", "Remaining Capex", "Commit?"] if m not in col]
    if missing:
        sys.exit(f"ERRO: colunas ausentes no cabecalho: {missing}")

    # primeira sequencia continua de colunas mensais (datas)
    months = []
    for j, h in enumerate(hdr):
        if isinstance(h, dt.datetime):
            if months and j != months[-1][0] + 1:
                break
            months.append((j, h.strftime("%Y-%m")))

    # data de corte: datetime na linha acima do cabecalho, nas colunas de metricas
    cutoff = None
    for j in range(min(len(rows[hi - 1]), col["Remaining Capex"] + 3)):
        v = rows[hi - 1][j]
        if isinstance(v, dt.datetime):
            cutoff = v.strftime("%Y-%m-%d")

    commit_col = col["Commit?"]
    projects, order, seen, block_rows, blocks = {}, [], {}, [], {}
    for i in range(hi + 1, len(rows)):
        r = rows[i] + [None] * (len(hdr) - len(rows[i]))
        a, b, tot = txt(r[0]), txt(r[1]), txt(r[commit_col])
        if tot == "Total":
            for k in block_rows:
                projects[k]["block"] = a
            cs = [projects[k]["country"] for k in block_rows]
            blocks[a] = {"block": a, "country": max(set(cs), key=cs.count) if cs else "(sem país)",
                         "m": {m: num(r[col[m]]) or 0.0 for m in TOTAL_METRICS if m in col}}
            block_rows = []
            continue
        if a and not b:  # fim da area de projetos (ex.: 'Expansions')
            break
        if not (a and b):
            continue
        if EXCLUDE_PROJECT.match(b):
            continue
        # linhas de reembolso de IVA ("VAT Refunded - ..."; Type 'Reimbursement') ficam fora da comparacao
        if "Type" in col and txt(r[col["Type"]]).lower() == "reimbursement":
            continue
        base = f"{a} | {b}"
        seen[base] = seen.get(base, 0) + 1
        key = base if seen[base] == 1 else f"{base} #{seen[base]}"
        country = txt(r[col["Country"]]) if "Country" in col else ""
        projects[key] = {
            "key": key, "site": a, "project": b, "row": i + 1, "block": "",
            "country": COUNTRY_FIX.get(country, country) or "(sem país)",
            "m": {m: num(r[col[m]]) if m in col else None for m in METRICS},
            "a": {k: txt(r[col[k]]) for k in ATTRS if k in col},
            "curve": [num(r[j]) or 0.0 for j, _ in months],
            "errors": sorted({txt(r[j]) for j in range(len(hdr)) if isinstance(r[j], str) and r[j].startswith("#")}),
        }
        order.append(key)
        block_rows.append(key)
    return {"file": os.path.basename(path), "cutoff": cutoff, "months": [m for _, m in months],
            "projects": projects, "order": order, "blocks": blocks}


# ---------------------------------------------------------------- comparacao
def compare(old, new, tol):
    months = new["months"]
    if old["months"] != months:  # alinhar curvas pela uniao de meses
        allm = sorted(set(old["months"]) | set(months))
        for s in (old, new):
            idx = {m: i for i, m in enumerate(s["months"])}
            for p in s["projects"].values():
                p["curve"] = [p["curve"][idx[m]] if m in idx else 0.0 for m in allm]
        months = allm
    cut = (new["cutoff"] or old["cutoff"] or "")[:7]

    O, N = old["projects"], new["projects"]
    added = [k for k in new["order"] if k not in O]
    removed = [k for k in old["order"] if k not in N]

    def pv(p):
        return {"key": p["key"], "site": p["site"], "project": p["project"], "country": p["country"],
                "block": p["block"], "row": p["row"], "m": p["m"], "a": p["a"]}

    # possiveis renomeacoes: mesmo site, mesmo bloco, um novo e um removido
    renames = []
    for r in removed:
        po = O[r]
        cands = [a for a in added if N[a]["site"] == po["site"] and N[a]["block"] == po["block"]]
        if cands:
            best = min(cands, key=lambda a: abs((N[a]["m"]["EFC"] or 0) - (po["m"]["EFC"] or 0)))
            renames.append({"old": r, "new": best, "site": po["site"]})

    changed, curve_changes = [], []
    for k in new["order"]:
        if k not in O:
            continue
        po, pn = O[k], N[k]
        deltas = {}
        for m in METRICS:
            a, b = po["m"][m], pn["m"][m]
            t = 0.0005 if m == "% Procured" else tol
            if (a is None) != (b is None) and (a or b):
                deltas[m] = {"old": a, "new": b, "d": (b or 0) - (a or 0)}
            elif a is not None and b is not None and abs(b - a) >= t:
                deltas[m] = {"old": a, "new": b, "d": b - a}
        attrs = {k2: {"old": po["a"].get(k2, ""), "new": pn["a"].get(k2, "")}
                 for k2 in ATTRS if po["a"].get(k2, "") != pn["a"].get(k2, "")}
        diff = [n - o for o, n in zip(po["curve"], pn["curve"])]
        mchg = [(months[i], po["curve"][i], pn["curve"][i], d) for i, d in enumerate(diff) if abs(d) >= tol]
        past = [x for x in mchg if cut and x[0] <= cut]
        if deltas or attrs:
            changed.append({**pv(pn), "deltas": deltas, "attrs": attrs,
                            "curve_months": len(mchg), "past_months": len(past)})
        if mchg:
            curve_changes.append({
                "key": k, "site": pn["site"], "project": pn["project"], "country": pn["country"],
                "n": len(mchg), "abs": sum(abs(x[3]) for x in mchg), "net": sum(x[3] for x in mchg),
                "past": [{"m": x[0], "old": x[1], "new": x[2], "d": x[3]} for x in past],
                "months": [{"m": x[0], "old": x[1], "new": x[2], "d": x[3]} for x in mchg],
            })
    curve_changes.sort(key=lambda x: -x["abs"])
    changed.sort(key=lambda x: -max([abs(v["d"]) for k2, v in x["deltas"].items() if k2 != "% Procured"] or [0]))

    # totais oficiais: linhas 'Total' de cada bloco da planilha (nao somamos as linhas de projeto,
    # porque as formulas de Total excluem algumas linhas, como 'Expansion').
    BO, BN = old["blocks"], new["blocks"]
    by_block = []
    for b in list(BN) + [b for b in BO if b not in BN]:
        a, n = BO.get(b), BN.get(b)
        row = {"block": b, "country": (n or a)["country"], "status": "" if a and n else ("novo" if n else "removido")}
        for m in TOTAL_METRICS:
            row[m] = {"old": a["m"].get(m, 0.0) if a else 0.0, "new": n["m"].get(m, 0.0) if n else 0.0}
        by_block.append(row)
    agg = {}
    for r in by_block:
        c = agg.setdefault(r["country"], {"country": r["country"], "blocks": 0,
                                          **{m: {"old": 0.0, "new": 0.0} for m in TOTAL_METRICS}})
        c["blocks"] += 1
        for m in TOTAL_METRICS:
            c[m]["old"] += r[m]["old"]; c[m]["new"] += r[m]["new"]
    for c in agg.values():
        c["n_old"] = sum(1 for p in O.values() if p["country"] == c["country"])
        c["n_new"] = sum(1 for p in N.values() if p["country"] == c["country"])
    by_country = sorted(agg.values(), key=lambda c: -c["EFC"]["new"])

    def curve_sum(s):
        tot = [0.0] * len(months)
        for p in s["projects"].values():
            for i, v in enumerate(p["curve"]):
                tot[i] += v
        return tot
    co, cn = curve_sum(old), curve_sum(new)
    # recortar meses sem valores nas duas versoes
    nz = [i for i in range(len(months)) if abs(co[i]) > 0.5 or abs(cn[i]) > 0.5]
    lo, hi = (nz[0], nz[-1] + 1) if nz else (0, 0)

    errs_new = [{"key": p["key"], "errors": p["errors"]} for p in N.values() if p["errors"]]

    # todas as linhas de projeto (uniao dos dois arquivos), para os filtros da pagina "Change Details"
    def dtype(p):
        cat, typ = p["a"].get("Capex categorization", ""), p["a"].get("Type", "")
        return "Spec - OFE" if cat == "OFE" else "Fitout" if cat == "Fitout" else (typ or "(vazio)")
    allp = []
    for k in list(new["order"]) + [k for k in old["order"] if k not in N]:
        ref = N.get(k) or O[k]
        po, pn = O.get(k), N.get(k)
        allp.append({
            "key": k, "p": f'{ref["site"]} {ref["project"]}', "country": ref["country"],
            "block": ref["block"],
            "other": ref["a"].get("Other", "") or "(vazio)", "type": dtype(ref),
            "cat": ref["a"].get("Capex categorization", "") or "(vazio)",
            "mo": {m: (po["m"][m] or 0.0) if po else 0.0 for m in TOTAL_METRICS},
            "mn": {m: (pn["m"][m] or 0.0) if pn else 0.0 for m in TOTAL_METRICS},
            "co": [round(v) for v in po["curve"][lo:hi]] if po else None,
            "cn": [round(v) for v in pn["curve"][lo:hi]] if pn else None,
        })

    return {
        "generated": dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
        "old": {"file": old["file"], "cutoff": old["cutoff"], "n": len(O)},
        "new": {"file": new["file"], "cutoff": new["cutoff"], "n": len(N)},
        "tol": tol, "cut_month": cut,
        "added": [pv(N[k]) for k in added],
        "removed": [pv(O[k]) for k in removed],
        "renames": renames,
        "changed": changed,
        "curve_changes": curve_changes,
        "by_country": by_country,
        "by_block": by_block,
        "curve": {"months": months[lo:hi], "old": co[lo:hi], "new": cn[lo:hi]},
        "errors_new": errs_new,
        "all": allp,
        "same_cutoff": old["cutoff"] == new["cutoff"],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--folder")
    ap.add_argument("--old")
    ap.add_argument("--new")
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--tol", type=float, default=1.0)
    a = ap.parse_args()
    if a.folder:
        a.old, a.new = pick_files(a.folder)
    if not (a.old and a.new):
        sys.exit("Informe --folder ou --old e --new")
    print(f"Base (Cons): {os.path.basename(a.old)}\nAtual (PROD): {os.path.basename(a.new)}", flush=True)
    old = read_summary(a.old)
    print(f"  lido base: {len(old['projects'])} projetos", flush=True)
    new = read_summary(a.new)
    print(f"  lido atual: {len(new['projects'])} projetos", flush=True)
    res = compare(old, new, a.tol)
    data = json.dumps(res, ensure_ascii=False, default=str)
    html = open(a.template, encoding="utf-8").read().replace("__DATA__", data.replace("</", "<\\/"))
    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    open(a.out, "w", encoding="utf-8").write(html)
    open(os.path.splitext(a.out)[0] + ".json", "w", encoding="utf-8").write(data)
    print(json.dumps({
        "novos": len(res["added"]), "removidos": len(res["removed"]), "renomeacoes": len(res["renames"]),
        "alterados": len(res["changed"]), "curva_alterada": len(res["curve_changes"]),
        "curva_meses_realizados": sum(1 for c in res["curve_changes"] if c["past"]),
        "corte_base": res["old"]["cutoff"], "corte_atual": res["new"]["cutoff"],
        "saida": a.out}, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
